import aiohttp
import asyncio
import os
import json
import pickle
import pandas as pd
from dotenv import load_dotenv
import openpyxl


load_dotenv()

from urllib3.exceptions import InsecureRequestWarning
import requests
requests.packages.urllib3.disable_warnings(InsecureRequestWarning)

def load_backup(file_name):
    if os.path.exists(file_name):
        with open(file_name, "rb") as f:
            return pickle.load(f)
    return None

def save_backup(file_name, processed_prodi, course_rows, fakultas_data):
    with open(file_name, "wb") as f:
        pickle.dump({
            "processed_prodi": processed_prodi,
            "course_rows": course_rows,
            "fakultas_data": fakultas_data
        }, f)

BASE_URL = "https://sikola-v2.unhas.ac.id"

JENIS_KELAS_MAP = {
    1: "Reguler",
    2: "Tanpa Pertemuan",
    3: "Kerjasama Dalam Negeri",
    4: "Internasional",
    5: "Penutup Starata"
}

async def get_sync_status(session, course_id):
    url = f"{BASE_URL}/grade/getSinkronNilai.php?courseId={course_id}"
    headers = {
        "Authorization": "Bearer ccd437b923b46aa49e922034903ac628a447f11f8cf4d464d7d21b790ccc6ab7e88a96fb5b4cee1a87ab1dddb3aea7e3a715ad9f0d9f730de09af7caa56b7a73044ef0a7be7fb15931152f0ecb66f78d1b202f1103d0820ccd445186896062f769a6818bbf1963a6790940b0059e47933637b60aecec366c7c62aebfce7f8e1f"
    }
    async with session.get(url, headers=headers) as response:
        data = await response.json()
        if data and len(data) > 0 and "status" in data[0]:
            return data[0]["status"] == 1
        return False

async def fetch_sikola_course():
    async with aiohttp.ClientSession() as session:
        active_semester = "TA242"
        backup_file = "NilaiPerKelas2-TA242.pkl"

        backup_data = load_backup(backup_file)
        processed_prodi = set(backup_data["processed_prodi"]) if backup_data and "processed_prodi" in backup_data else set()
        course_rows = backup_data["course_rows"] if backup_data and "course_rows" in backup_data else []
        fakultas_data = backup_data["fakultas_data"] if backup_data and "fakultas_data" in backup_data else {}

        with open("data/prodi_semester.json", "r", encoding="utf-8") as faculty_file:
            datafakultasProdi = json.load(faculty_file)

        for prodi in datafakultasProdi["prodis"]:
            nama_prodi = prodi["nama_resmi"]
            if nama_prodi in processed_prodi:
                continue

            fakultas = prodi["fakultas"]["nama_resmi"]

            param_prodi = {
                "wsfunction": "core_course_get_categories",
                "criteria[0][key]": "name",
                "criteria[0][value]": nama_prodi,
            }
            response = await session.get(
                BASE_URL + "/webservice/rest/server.php?wstoken=3bdff39b5f62204b247c93e852caad8b&moodlewsrestformat=json",
                params=param_prodi, ssl=False
            )
            response = await response.json()
            if not response:
                continue

            param_prodi_courses = {
                "wsfunction": "core_course_get_courses_by_field",
                "field": "category",
                "value": response[0]["id"],
            }
            response = await session.get(
                BASE_URL + "/webservice/rest/server.php?wstoken=3bdff39b5f62204b247c93e852caad8b&moodlewsrestformat=json",
                params=param_prodi_courses, ssl=False
            )
            result_course = await response.json()
            if "courses" not in result_course:
                continue

            for course in result_course["courses"]:
                if active_semester in course["shortname"]:
                    if '.' in course['idnumber']:
                        id_kelas = course['idnumber'].split(".")[1]
                        detail_kelas_path = f"data/detailKelas/TA242.10/{id_kelas}.json"
                        if os.path.exists(detail_kelas_path):
                            with open(detail_kelas_path, "r", encoding="utf-8") as kelas_file:
                                isiKelas = json.load(kelas_file)
                                id_kelas_kuliah_jenis = isiKelas.get('id_kelas_kuliah_jenis', None)
                                jenis_kelas = JENIS_KELAS_MAP.get(id_kelas_kuliah_jenis, "Lainnya")
                                sync_status = await get_sync_status(session, course["id"])
                                course_rows.append({
                                    "Nama Course": course.get("fullname", ""),
                                    "Status Sinkron": "Sinkron" if sync_status else "Tidak Sinkron",
                                    # "id_kelas_kuliah_jenis": id_kelas_kuliah_jenis,
                                    "Jenis Kelas": jenis_kelas ,
                                    "id_kelas_kuliah": id_kelas,
                                    "Nama Prodi": nama_prodi,
                                    "Nama Fakultas": fakultas,
                                    "Link Sikola": f"{BASE_URL}/grade/report/grader/index.php?id={course['id']}",
                                })
                                # Simpan per fakultas untuk sheet
                                if fakultas not in fakultas_data:
                                    fakultas_data[fakultas] = []
                                fakultas_data[fakultas].append(course_rows[-1])

            processed_prodi.add(nama_prodi)
            print(nama_prodi)
            save_backup(backup_file, processed_prodi, course_rows, fakultas_data)

        # Tulis ke file Excel: sheet per fakultas, baris per course
        with pd.ExcelWriter("course_per_fakultas_per_course_TA242_NilaiPerKelas2.xlsx") as writer:
            for fakultas, rows in fakultas_data.items():
                df = pd.DataFrame(rows)
                df.to_excel(writer, sheet_name=fakultas[:31], index=False)
            
        wb = openpyxl.load_workbook("course_per_fakultas_per_course_TA242_NilaiPerKelas2.xlsx")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # Cari kolom Link Sikola
            for col in ws.iter_cols(1, ws.max_column):
                if col[0].value == "Link Sikola":
                    link_col_idx = col[0].column
                    for cell in col[1:]:
                        if cell.value:
                            cell.hyperlink = cell.value
                            cell.style = "Hyperlink"
        wb.save("course_per_fakultas_per_course_TA242_NilaiPerKelas2.xlsx")


        if os.path.exists(backup_file):
            os.remove(backup_file)

if __name__ == "__main__":
    asyncio.run(fetch_sikola_course())